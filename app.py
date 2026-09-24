import io
import uuid
from datetime import datetime

import openpyxl
import pandas as pd
from flask import Flask, render_template, request, send_file, flash, redirect, url_for, abort

import engine
import correlativo
import salida_sap

app = Flask(__name__)
app.secret_key = "cambiar-esta-clave-en-produccion"

# Archivos generados listos para descargar, en memoria (proceso único de Flask
# dev server — se pierden si se reinicia, es intencional: son de un solo uso).
# Evita el patrón "flash + send_file directo": si /procesar devolviera el
# archivo de una, los avisos quedarían pegados en la sesión sin mostrarse
# (send_file no renderiza plantilla) y reaparecerían solos en cualquier
# recarga posterior de página, sin relación con lo que el usuario hizo.
# Con Post/Redirect/Get, los avisos se muestran y se consumen una sola vez,
# en la página de resultado, y la descarga es un segundo click aparte.
_DESCARGAS_PENDIENTES: dict[str, tuple[bytes, str, str]] = {}


@app.route("/", methods=["GET"])
def index():
    # Página intermedia: elegir primero Repuestos o Modelo de Unidades (son
    # procesos distintos, con inputs distintos) y recién después subir el
    # archivo, para que nadie cargue un input en el flujo equivocado.
    tipos = engine.listar_tipos()
    return render_template("seleccion.html", tipos=tipos)


@app.route("/ampliar/<tipo_id>", methods=["GET"])
def ampliar(tipo_id):
    try:
        cfg = engine.cargar_config(tipo_id)
    except engine.TipoNoEncontradoError as e:
        flash(str(e), "error")
        return redirect(url_for("index"))
    tipo = {
        "id": cfg["id"],
        "nombre": cfg["nombre"],
        "descripcion": cfg.get("descripcion", ""),
        "tiene_diccionario": bool(cfg.get("diccionario_referencia")),
    }
    return render_template("index.html", tipo=tipo)


@app.route("/plantilla/<tipo_id>")
def plantilla(tipo_id):
    try:
        wb = engine.generar_plantilla_vacia(tipo_id)
    except engine.TipoNoEncontradoError as e:
        flash(str(e), "error")
        return redirect(url_for("index"))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"plantilla_{tipo_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/diccionario/<tipo_id>")
def diccionario(tipo_id):
    try:
        cfg = engine.cargar_config(tipo_id)
    except engine.TipoNoEncontradoError as e:
        flash(str(e), "error")
        return redirect(url_for("index"))

    diccionario_cfg = cfg.get("diccionario_referencia")
    diccionarios_cfg = cfg.get("diccionarios_referencia")
    if not diccionario_cfg and not diccionarios_cfg:
        flash(f"El tipo '{tipo_id}' no tiene un diccionario de referencia configurado.", "error")
        return redirect(url_for("index"))

    if diccionario_cfg:
        origen_path = engine.BASE_DIR / diccionario_cfg["reference_file"]
        if not origen_path.exists():
            flash(f"Falta el archivo de diccionario: {origen_path}", "error")
            return redirect(url_for("index"))
        return send_file(
            origen_path,
            as_attachment=True,
            download_name=origen_path.name,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # Varios diccionarios (Repuestos): se bundlean en un solo Excel, una hoja
    # por diccionario, igual que se hace dentro de la plantilla descargable.
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for d_cfg in diccionarios_cfg:
        engine._agregar_hoja_diccionario(wb, d_cfg)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"diccionarios_{tipo_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/procesar", methods=["POST"])
def procesar():
    tipo_id = request.form.get("tipo")
    archivo = request.files.get("archivo")
    generar_sap = request.form.get("generar_sap") == "on"  # Checkbox para habilitar SAP

    if not tipo_id:
        flash("Selecciona un tipo de ampliación.", "error")
        return redirect(url_for("index"))

    volver = redirect(url_for("ampliar", tipo_id=tipo_id))

    if not archivo or archivo.filename == "":
        flash("Sube un archivo Excel con el input.", "error")
        return volver

    try:
        resultado_df = engine.procesar(tipo_id, archivo)
    except (engine.TipoNoEncontradoError, engine.InputInvalidoError) as e:
        flash(str(e), "error")
        return volver

    avisos = resultado_df.attrs.get("avisos", [])

    # Si se solicita generar formato SAP
    df_salida = resultado_df
    nombre_sheet = "AMPLIADO"
    df_pendientes = None

    if generar_sap:
        if "FILIAL CODIGO" not in resultado_df.columns:
            flash("⚠️ Columna FILIAL CODIGO no encontrada. No se puede generar SAP.", "warning")
            generar_sap = False
        else:
            filiales_unicas = resultado_df["FILIAL CODIGO"].unique()
            if len(filiales_unicas) > 1:
                flash("⚠️ Entrada con múltiples filiales en el mismo archivo. Solo se puede generar SAP para una filial a la vez.", "error")
                return volver
            filial = str(filiales_unicas[0]).strip()

            try:
                # Repuestos: "Serie o Lote?" define ZRP1/ZRP2/ZRP3 (columna
                # "TIPO MATERIAL REPUESTO", ya traducida por engine.py).
                # Modelo de Unidades: la columna "TIPO MATERIAL" del input
                # (ZVEH/ZMAQ/ZCAM/ZUSA) define plantilla y rango. En ambos
                # casos puede haber una mezcla de tipos en el mismo archivo,
                # y cada tipo sale en su propia hoja.
                if tipo_id == "repuestos":
                    col_tipo_sap = "TIPO MATERIAL REPUESTO"
                    tipos_sap = ("ZRP1", "ZRP2", "ZRP3")
                else:
                    col_tipo_sap = "TIPO MATERIAL"
                    tipos_sap = tuple(dict.fromkeys(resultado_df[col_tipo_sap]))
                partes_sap_por_tipo = {}
                pendientes = {}
                obligatorios_vacios = []
                marcas_sin_categoria = set()
                marcas_sin_grupo_compras = set()
                npf_largos = []
                npf_duplicados = []
                for tipo_material_sap in tipos_sap:
                    subset = resultado_df[resultado_df[col_tipo_sap] == tipo_material_sap]
                    if subset.empty:
                        continue
                    if tipo_id == "repuestos" and tipo_material_sap != "ZRP1":
                        materiales = subset.drop_duplicates("NUMERO MATERIAL")["TEXTO BREVE"].tolist()
                        avisos.append(
                            f"🔵 {tipo_material_sap} ({len(materiales)} material(es)): " + ", ".join(materiales)
                        )
                    df_parte, meta_parte = salida_sap.aplicar_plantilla_sap(subset, tipo_material_sap)
                    partes_sap_por_tipo[tipo_material_sap] = df_parte
                    pendientes.update(meta_parte.get("campos_pendientes", {}))
                    obligatorios_vacios.extend(meta_parte.get("columnas_obligatorias_vacias", []))
                    marcas_sin_categoria.update(meta_parte.get("marcas_sin_categoria_valoracion", []))
                    marcas_sin_grupo_compras.update(meta_parte.get("marcas_sin_grupo_compras", []))
                    npf_largos.extend(meta_parte.get("npf_largos", []))
                    npf_duplicados.extend(meta_parte.get("npf_duplicados", []))

                if npf_largos:
                    avisos.append(
                        f"🔵 NPF de más de 18 caracteres ({len(npf_largos)}): se usó el correlativo en lugar del NPF "
                        f"en Nº antiguo material / Número de artículo Europeo: " + ", ".join(dict.fromkeys(npf_largos))
                    )
                if npf_duplicados:
                    avisos.append(
                        f"🔵 NPF:Fabricante ya existente ({len(npf_duplicados)}): se dejó el correlativo en el EAN "
                        f"(Número de artículo Europeo): " + ", ".join(dict.fromkeys(npf_duplicados))
                    )

                if marcas_sin_categoria:
                    avisos.append(
                        f"⚠️ Sin Categoría valoración confirmada para marca(s): "
                        f"{', '.join(sorted(marcas_sin_categoria))} (ver docs/categoria_valoracion_pendientes.md)."
                    )

                if marcas_sin_grupo_compras:
                    avisos.append(
                        f"⚠️ Sin Grupo de compras confirmado para marca(s): "
                        f"{', '.join(sorted(marcas_sin_grupo_compras))} (ver config/grupo_compras.json)."
                    )

                if pendientes:
                    avisos.append(f"ℹ️ {len(pendientes)} columnas pendientes de negocio (ver hoja PENDIENTES).")
                    df_pendientes = pd.DataFrame(
                        {"Columna SAP pendiente": list(pendientes.keys()), "Motivo / nota del negocio": list(pendientes.values())}
                    )

                faltantes_no_pendientes = [c for c in dict.fromkeys(obligatorios_vacios) if c not in pendientes]
                if faltantes_no_pendientes:
                    avisos.append(
                        f"⚠️ Sin dato (no pendiente de negocio, revisar tabla de referencia): "
                        f"{', '.join(faltantes_no_pendientes)}"
                    )
            except salida_sap.FormatoSAPError as e:
                flash(f"Error en conversión SAP: {e}", "error")
                return volver

    for a in avisos:
        flash(a, "warning")

    buffer = io.BytesIO()
    if generar_sap and tipo_id == "repuestos":
        # El programa que carga esto a SAP espera el mismo layout que
        # PlanillaCargaTattersall_Repuestos_ouput.xlsx (encabezados hasta la
        # fila 5, datos desde la fila 6) — una hoja por tipo de material
        # (ZRP1/ZRP2/ZRP3), no todo junto en una sola hoja simple.
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for tipo_material_sap, df_parte in partes_sap_por_tipo.items():
            salida_sap.escribir_hoja_sap_repuestos(wb, tipo_material_sap, df_parte)
        if df_pendientes is not None:
            ws_pend = wb.create_sheet(title="PENDIENTES")
            ws_pend.append(list(df_pendientes.columns))
            for fila in df_pendientes.itertuples(index=False):
                ws_pend.append(list(fila))
        wb.save(buffer)
    else:
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            if generar_sap:
                for tipo_material_sap, df_parte in partes_sap_por_tipo.items():
                    df_parte.to_excel(writer, index=False, sheet_name=f"SAP_{tipo_material_sap}")
            else:
                df_salida.to_excel(writer, index=False, sheet_name=nombre_sheet)
            if df_pendientes is not None:
                df_pendientes.to_excel(writer, index=False, sheet_name="PENDIENTES")

    tipo_sufijo = "SAP" if generar_sap else "ampliado"
    nombre_salida = f"{tipo_id}_{tipo_sufijo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    token = uuid.uuid4().hex
    _DESCARGAS_PENDIENTES[token] = (buffer.getvalue(), nombre_salida, tipo_id)

    return redirect(url_for("resultado", token=token))


@app.route("/resultado/<token>")
def resultado(token):
    if token not in _DESCARGAS_PENDIENTES:
        flash("El archivo generado ya no está disponible (probablemente ya lo descargaste). Genera la ampliación de nuevo.", "error")
        return redirect(url_for("index"))
    _, nombre_salida, tipo_id = _DESCARGAS_PENDIENTES[token]
    return render_template("resultado.html", token=token, nombre_salida=nombre_salida, tipo_id=tipo_id)


@app.route("/descargar/<token>")
def descargar(token):
    # Un solo uso: se saca del diccionario apenas se sirve, para no acumular
    # archivos en memoria indefinidamente.
    entrada = _DESCARGAS_PENDIENTES.pop(token, None)
    if entrada is None:
        abort(404)
    contenido, nombre_salida, _ = entrada

    return send_file(
        io.BytesIO(contenido),
        as_attachment=True,
        download_name=nombre_salida,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/historial")
def historial():
    tipos_cfg = {t["id"]: engine.cargar_config(t["id"]) for t in engine.listar_tipos()}
    estados = []
    for tipo_id, cfg in tipos_cfg.items():
        corr_cfg = cfg.get("correlativo", {})
        if corr_cfg.get("enabled"):
            try:
                estado = correlativo.estado_contador(
                    corr_cfg["nombre"], corr_cfg.get("range_min"), corr_cfg.get("range_max")
                )
                estados.append((tipo_id, estado))
            except Exception:
                pass  # Skip si no se puede obtener estado
    
    registros = correlativo.historial(limit=200)
    return render_template("historial.html", estados=estados, registros=registros)


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)

